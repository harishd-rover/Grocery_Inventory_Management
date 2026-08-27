from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

OUT = 'grocery_inventory_explanation.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(247, 249, 246)
INK = RGBColor(35, 50, 45)
MUTED = RGBColor(103, 119, 110)
GREEN = RGBColor(38, 115, 84)
PALE_GREEN = RGBColor(226, 242, 232)
YELLOW = RGBColor(210, 155, 39)
PALE_YELLOW = RGBColor(255, 245, 218)
BLUE = RGBColor(67, 119, 157)
PALE_BLUE = RGBColor(229, 240, 248)
RED = RGBColor(198, 88, 77)
PALE_RED = RGBColor(252, 233, 229)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(28, 43, 37)

FONT = 'Aptos'
MONO = 'Consolas'


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.06):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_text(slide, runs, x, y, w, h, size=16, color=INK, font=FONT, margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    for item in runs:
        r = p.add_run(); r.text = item[0]
        r.font.name = item[2] if len(item) > 2 else font
        r.font.size = Pt(item[1] if len(item) > 1 else size)
        r.font.bold = item[3] if len(item) > 3 else False
        r.font.color.rgb = item[4] if len(item) > 4 else color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=RGBColor(225, 234, 227), radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color; shape.line.width = Pt(width)
    return shape


def bullet_box(slide, items, x, y, w, h, size=16, color=INK, fill=WHITE):
    rect(slide, x, y, w, h, fill, radius=True)
    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.14), Inches(w - 0.36), Inches(h - 0.28))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = FONT; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(8)
        p.bullet = True
    return box


def title(slide, kicker, heading, number):
    textbox(slide, kicker.upper(), 0.55, 0.35, 3.7, 0.28, 10, GREEN, True)
    textbox(slide, heading, 0.55, 0.7, 11.9, 0.65, 28, INK, True)
    textbox(slide, f'{number:02d}', 12.25, 0.38, 0.55, 0.32, 11, MUTED, True, align=PP_ALIGN.RIGHT)
    line(slide, 0.55, 1.48, 12.78, 1.48, RGBColor(220, 230, 223), 1)


def footer(slide, text='Grocery Inventory Management'):
    textbox(slide, text, 0.58, 7.16, 5.5, 0.18, 8, MUTED)


def add_slide(kicker, heading, number, bg=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide, bg)
    title(slide, kicker, heading, number); footer(slide)
    return slide


def code_box(slide, code, x, y, w, h, size=12):
    rect(slide, x, y, w, h, DARK, DARK, radius=True)
    textbox(slide, code, x + 0.16, y + 0.12, w - 0.32, h - 0.24, size, WHITE, False, MONO)


def note(slide, text):
    # Add a small explanation strip that makes slides usable without a presenter.
    rect(slide, 0.58, 6.55, 12.15, 0.42, PALE_GREEN, PALE_GREEN, radius=True)
    textbox(slide, 'KEY IDEA  ' + text, 0.78, 6.64, 11.75, 0.2, 10, GREEN, True)

# 1
slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide, DARK)
textbox(slide, 'GROCERY INVENTORY MANAGEMENT', 0.75, 1.05, 8.7, 0.35, 13, RGBColor(143, 210, 170), True)
textbox(slide, 'Frontend and full-stack\nproject explained', 0.72, 1.6, 9.8, 1.65, 38, WHITE, True)
textbox(slide, 'A code-grounded walkthrough of the Vite frontend, Flask API,\nauthentication, workflows, and current architecture.', 0.78, 3.65, 8.4, 0.7, 17, RGBColor(211, 226, 216))
rect(slide, 9.65, 1.2, 2.25, 3.8, GREEN, GREEN, radius=True)
textbox(slide, 'g', 10.25, 1.65, 1.05, 1.2, 78, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
textbox(slide, 'grocerly.', 9.95, 3.2, 1.65, 0.4, 18, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 'PROJECT GUIDE', 0.78, 6.55, 3, 0.25, 10, RGBColor(143, 210, 170), True)
textbox(slide, '2026', 11.8, 6.55, 0.75, 0.25, 10, RGBColor(211, 226, 216), True, align=PP_ALIGN.RIGHT)

# 2
slide = add_slide('Orientation', 'What this project contains', 2)
bullet_box(slide, ['Browser application: Vite serves a single HTML entry point.', 'UI implementation: vanilla JavaScript builds HTML strings and binds events.', 'Backend: Flask exposes JSON REST endpoints under /api.', 'Storage: Python in-memory lists and dictionaries for development.', 'Users: bearer-token login with admin and staff permissions.'], 0.65, 1.85, 5.7, 4.35, 17)
rect(slide, 6.8, 1.95, 5.35, 3.95, PALE_BLUE, PALE_BLUE, radius=True)
textbox(slide, 'The central idea', 7.15, 2.25, 4.5, 0.35, 20, BLUE, True)
textbox(slide, 'The frontend is a small single-page application. It keeps the current screen and loaded records in one state object, then re-renders the visible page whenever state changes.', 7.15, 2.8, 4.45, 1.4, 17, INK)
textbox(slide, 'Browser  <->  Flask API  <->  In-memory records', 7.15, 4.7, 4.55, 0.65, 15, DARK, True, MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
note(slide, 'The system is intentionally simple enough to understand, but already separates UI and API responsibilities.')

# 3
slide = add_slide('Project map', 'How the folders fit together', 3)
items = [('frontend/index.html', 'HTML shell\n#root mount point', PALE_BLUE, BLUE), ('frontend/src/main.js', 'Application logic\nstate, views, API, events', PALE_GREEN, GREEN), ('frontend/src/styles.css', 'Visual system\nlayout, tables, mobile CSS', PALE_YELLOW, YELLOW), ('backend/app.py', 'REST API\nauth, roles, business rules', PALE_RED, RED), ('database/seed_data.json', 'Reference data\nnot loaded at startup', RGBColor(239, 235, 249), RGBColor(111, 82, 145))]
for i, (name, desc, fill, accent) in enumerate(items):
    x = 0.75 + (i % 3) * 4.15; y = 1.9 + (i // 3) * 2.2
    rect(slide, x, y, 3.55, 1.55, fill, fill, radius=True)
    textbox(slide, name, x + 0.2, y + 0.22, 3.1, 0.3, 14, accent, True, MONO)
    textbox(slide, desc, x + 0.2, y + 0.68, 3.05, 0.6, 15, INK)
line(slide, 4.3, 3.45, 5.1, 3.45, GREEN, 2); line(slide, 8.45, 3.45, 9.25, 3.45, GREEN, 2)
textbox(slide, 'Browser UI', 4.45, 3.05, 1.1, 0.25, 10, MUTED, True, align=PP_ALIGN.CENTER)
textbox(slide, 'API boundary', 8.55, 3.05, 1.1, 0.25, 10, MUTED, True, align=PP_ALIGN.CENTER)
note(slide, 'index.html starts the app; main.js controls behavior; app.py is the trusted business boundary.')

# 4
slide = add_slide('Startup', 'What happens when the frontend opens', 4)
steps = [('1', 'index.html', 'Creates <div id="root"> and loads main.js.'), ('2', 'main.js', 'Imports styles.css, reads localStorage, initializes state.'), ('3', 'render()', 'Shows login if no user exists; otherwise shows the app shell.'), ('4', 'loadData()', 'Requests summary, products, suppliers, purchases, and admin data.'), ('5', 'bindEvents()', 'Connects buttons, forms, navigation, search, and modal actions.')]
for i, (n, head, desc) in enumerate(steps):
    y = 1.8 + i * 0.88
    rect(slide, 0.85, y, 0.52, 0.52, GREEN, GREEN, radius=True)
    textbox(slide, n, 0.85, y + 0.08, 0.52, 0.28, 14, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, head, 1.65, y - 0.02, 2.5, 0.28, 16, GREEN, True, MONO)
    textbox(slide, desc, 4.15, y - 0.02, 7.7, 0.42, 15, INK)
    if i < 4: line(slide, 1.11, y + 0.52, 1.11, y + 0.88, RGBColor(177, 205, 188), 1.5)
code_box(slide, "function render() {\n  root.innerHTML = state.user\n    ? renderApp()\n    : renderLogin()\n  bindEvents()\n}", 8.75, 5.85, 3.75, 0.7, 10)
note(slide, 'The app is a re-rendering UI: state changes cause HTML to be regenerated, then listeners are attached again.')

# 5
slide = add_slide('Frontend core', 'The state object is the UI memory', 5)
code_box(slide, "const state = {\n  user: readUser(),\n  page: 'Overview',\n  summary: null,\n  products: [], suppliers: [],\n  purchases: [], sales: [], users: [],\n  search: '', modal: null,\n  loading: false, billing: {}\n}", 0.75, 1.85, 5.1, 3.8, 14)
bullet_box(slide, ['user: who is signed in', 'page: current section', 'arrays: records loaded from Flask', 'modal: which dialog is open', 'billing: selected products and quantities', 'loading: whether requests are in progress'], 6.25, 1.9, 5.7, 3.7, 16, INK, PALE_GREEN)
note(slide, 'There is no framework store. The plain JavaScript object is the single source of truth for the current UI.')

# 6
slide = add_slide('Rendering', 'How screens are produced', 6)
bullet_box(slide, ['renderLogin() creates the sign-in form.', 'renderApp() creates sidebar, header, page content, and optional modal.', 'renderPage() chooses a page-specific renderer.', 'renderOverview(), renderDirectory(), renderBilling(), renderReports(), and renderUsers() return HTML.', 'escapeHtml() protects values inserted into HTML tables and labels.'], 0.7, 1.82, 5.7, 4.2, 16)
rect(slide, 6.8, 1.95, 5.25, 3.9, WHITE, RGBColor(225, 234, 227), radius=True)
textbox(slide, 'Page selection', 7.15, 2.25, 4.5, 0.3, 18, GREEN, True)
textbox(slide, 'state.page', 7.15, 2.9, 1.65, 0.38, 17, WHITE, True, MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line(slide, 8.9, 3.1, 9.45, 3.1, GREEN, 2)
textbox(slide, 'renderPage()', 9.55, 2.9, 1.95, 0.38, 15, WHITE, True, MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line(slide, 10.55, 3.45, 10.55, 3.88, GREEN, 2)
textbox(slide, 'Overview | Inventory | Billing\nReports | Users | directories', 8.0, 4.0, 5.0, 0.75, 16, INK, True, align=PP_ALIGN.CENTER)
code_box(slide, "state.page = 'Inventory'\nrender()", 7.55, 5.08, 2.7, 0.62, 11)
note(slide, 'Navigation is controlled by changing state.page, then rendering the new page.')

# 7
slide = add_slide('Authentication', 'Login and token flow', 7)
flow = [('Login form', PALE_BLUE, BLUE), ('POST /auth/login', PALE_GREEN, GREEN), ('Flask validates', PALE_YELLOW, YELLOW), ('Token + user', PALE_GREEN, GREEN), ('localStorage', PALE_BLUE, BLUE)]
for i, (label, fill, accent) in enumerate(flow):
    x = 0.7 + i * 2.45
    rect(slide, x, 2.1, 1.95, 1.05, fill, fill, radius=True)
    textbox(slide, label, x + 0.08, 2.35, 1.79, 0.45, 14, accent, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i < 4: line(slide, x + 1.95, 2.62, x + 2.38, 2.62, MUTED, 1.5)
code_box(slide, "localStorage.setItem('grocery_token', result.token)\nlocalStorage.setItem('grocery_user', JSON.stringify(result.user))", 1.1, 4.15, 6.15, 0.85, 12)
bullet_box(slide, ['apiFetch() reads the token and adds Authorization: Bearer <token>.', 'A 401 response calls logout().', 'Passwords are never stored in the frontend; Flask stores password hashes.'], 7.65, 3.85, 4.65, 2.0, 15, INK, PALE_RED)
note(slide, 'The token identifies the session; every protected API request depends on it.')

# 8
slide = add_slide('Authorization', 'Admin and staff roles', 8)
rect(slide, 0.75, 1.85, 5.65, 3.95, PALE_GREEN, PALE_GREEN, radius=True)
textbox(slide, 'Administrator', 1.1, 2.15, 4.7, 0.4, 22, GREEN, True)
textbox(slide, 'Can access:', 1.1, 2.78, 1.7, 0.3, 14, INK, True)
textbox(slide, 'Overview, Inventory, Purchases, Suppliers, Billing, Sales, Reports, Users\n\nCan add products, suppliers, purchases, bills, and users.', 1.1, 3.18, 4.7, 1.35, 16, INK)
rect(slide, 6.9, 1.85, 5.65, 3.95, PALE_BLUE, PALE_BLUE, radius=True)
textbox(slide, 'Staff member', 7.25, 2.15, 4.7, 0.4, 22, BLUE, True)
textbox(slide, 'Can access:', 7.25, 2.78, 1.7, 0.3, 14, INK, True)
textbox(slide, 'Inventory and Billing\n\nCan create sales and purchases, but cannot manage users or master data.', 7.25, 3.18, 4.7, 1.35, 16, INK)
code_box(slide, "@app.post('/api/users')\n@require_auth({'admin'})", 4.1, 5.95, 4.9, 0.62, 12)
note(slide, 'The frontend hides controls for staff, and the backend independently rejects unauthorized requests with HTTP 403.')

# 9
slide = add_slide('API client', 'How the frontend talks to Flask', 9)
code_box(slide, "function apiFetch(path, options = {}) {\n  const token = localStorage.getItem('grocery_token')\n  const headers = { ... }\n  if (token) headers.Authorization = `Bearer ${token}`\n  return fetch(`${API}${path}`, { ...options, headers })\n}", 0.75, 1.85, 6.0, 2.8, 13)
rect(slide, 7.25, 1.9, 4.9, 3.55, WHITE, RGBColor(225, 234, 227), radius=True)
textbox(slide, 'apiJson(path, options)', 7.65, 2.25, 4.1, 0.35, 19, GREEN, True, MONO)
textbox(slide, '1. Calls apiFetch()\n2. Parses JSON\n3. Checks response.ok\n4. Throws data.error on failure\n5. Returns usable data', 7.65, 2.9, 3.9, 1.8, 17, INK)
textbox(slide, 'API base URL: http://localhost:5000/api', 1.0, 5.25, 5.3, 0.35, 15, BLUE, True, MONO)
note(slide, 'Centralizing requests keeps authentication headers and error handling consistent across every feature.')

# 10
slide = add_slide('Data loading', 'Dashboard loading and refresh behavior', 10)
flow = [('loadData()', PALE_BLUE, BLUE), ('Promise.allSettled()', PALE_GREEN, GREEN), ('summary + records', PALE_YELLOW, YELLOW), ('state updated', PALE_GREEN, GREEN), ('render()', PALE_BLUE, BLUE)]
for i, (label, fill, accent) in enumerate(flow):
    x = 0.75 + (i % 5) * 2.45
    rect(slide, x, 2.05, 1.95, 1.08, fill, fill, radius=True)
    textbox(slide, label, x + 0.05, 2.31, 1.85, 0.48, 14, accent, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i < 4: line(slide, x + 1.95, 2.59, x + 2.4, 2.59, MUTED, 1.5)
code_box(slide, "const calls = [\n  apiJson('/summary'), apiJson('/products'),\n  apiJson('/suppliers'), apiJson('/purchases'),\n  admin ? apiJson('/sales') : Promise.resolve([]),\n  admin ? apiJson('/users') : Promise.resolve([])\n]", 1.0, 4.25, 6.2, 1.45, 12)
bullet_box(slide, ['loading-bar appears while requests are active', 'Promise.allSettled prevents one failed call from stopping every view', 'successful records replace old state before the next render'], 7.65, 4.25, 4.65, 1.55, 14, INK, PALE_GREEN)
note(slide, 'After adding a product, purchase, sale, or user, the frontend calls loadData() again to refresh the interface.')

# 11
slide = add_slide('User workflow', 'Why adding a user works only for admin', 11)
steps = [('Sign in as nishi', 'role = admin'), ('Open Users', 'Users navigation is visible'), ('Click Add user', 'state.modal = user'), ('Submit form', 'POST /api/users'), ('Backend checks role', 'admin accepted'), ('Refresh list', 'loadData()')]
for i, (a, b) in enumerate(steps):
    x = 0.65 + (i % 3) * 4.15; y = 1.9 + (i // 3) * 2.1
    rect(slide, x, y, 3.55, 1.28, PALE_GREEN if i != 4 else PALE_YELLOW, PALE_GREEN, radius=True)
    textbox(slide, a, x + 0.18, y + 0.2, 3.15, 0.3, 15, GREEN, True)
    textbox(slide, b, x + 0.18, y + 0.65, 3.15, 0.28, 14, INK, False, MONO)
    if i in (0, 1, 3): line(slide, x + 3.55, y + 0.64, x + 4.05, y + 0.64, MUTED, 1.3)
code_box(slide, "const routes = { user: '/users' }\nawait apiJson(routes.user, {\n  method: 'POST', body: JSON.stringify(data)\n})", 4.0, 6.0, 5.25, 0.62, 11)
note(slide, 'A staff login cannot add users by design: the API returns 403 Admin access required.')

# 12
slide = add_slide('Business workflows', 'Inventory, purchases, billing, and reports', 12)
items = [('Inventory', 'GET /products\nLow-stock check\nQuantity + price', PALE_GREEN, GREEN), ('Purchase', 'POST /purchases\nIncrease stock\nRecord supplier', PALE_YELLOW, YELLOW), ('Billing', 'POST /sales\nDecrease stock\nCreate invoice', PALE_BLUE, BLUE), ('Reports', 'GET /summary\nSales - purchases\nStock health', PALE_RED, RED)]
for i, (head, body, fill, accent) in enumerate(items):
    x = 0.7 + i * 3.1
    rect(slide, x, 2.0, 2.7, 2.45, fill, fill, radius=True)
    textbox(slide, head, x + 0.18, 2.3, 2.32, 0.35, 20, accent, True, align=PP_ALIGN.CENTER)
    textbox(slide, body, x + 0.25, 3.0, 2.2, 1.0, 16, INK, align=PP_ALIGN.CENTER)
line(slide, 3.38, 5.05, 9.95, 5.05, GREEN, 2)
textbox(slide, 'Products are the shared record: purchases raise quantity; sales lower quantity; reports summarize both.', 1.2, 5.35, 10.8, 0.52, 17, INK, True, align=PP_ALIGN.CENTER)
note(slide, 'The frontend collects input; Flask applies stock changes and validates the transaction.')

# 13
slide = add_slide('Backend', 'What Flask is responsible for', 13)
bullet_box(slide, ['Authentication: validates credentials and creates active bearer tokens.', 'Authorization: require_auth() checks token and allowed role.', 'Validation: required fields, duplicate usernames, product and supplier existence.', 'Business rules: stock cannot go below zero; at least one admin must remain.', 'Responses: JSON data plus meaningful HTTP status codes such as 201, 400, 401, 403, 404, and 409.'], 0.7, 1.85, 6.1, 4.35, 16)
rect(slide, 7.25, 1.95, 4.75, 3.9, DARK, DARK, radius=True)
textbox(slide, 'REST endpoint families', 7.6, 2.25, 4.05, 0.3, 18, RGBColor(143, 210, 170), True)
textbox(slide, '/auth/login\n/auth/me\n/users\n/summary\n/products\n/sales\n/suppliers\n/purchases\n/categories', 7.6, 2.85, 3.9, 2.35, 17, WHITE, False, MONO)
note(slide, 'Business rules belong in the backend because frontend visibility is not a security boundary.')

# 14
slide = add_slide('Styling', 'How the visual interface is organized', 14)
bullet_box(slide, ['CSS variables define ink, muted text, borders, green, yellow, red, and blue tones.', 'The app shell uses a fixed sidebar and scrollable main content.', 'Reusable visual classes cover panels, buttons, tables, metrics, alerts, forms, and modals.', 'DM Sans is used for body text; Manrope is used for headings and totals.', 'Media queries at 1050px and 700px adapt navigation, grids, forms, and reports.'], 0.7, 1.83, 6.0, 4.3, 16)
rect(slide, 7.25, 2.0, 4.85, 3.55, WHITE, RGBColor(225, 234, 227), radius=True)
textbox(slide, 'Responsive behavior', 7.6, 2.3, 4.1, 0.35, 19, GREEN, True)
textbox(slide, 'Desktop\nSidebar visible\nMulti-column metrics\n\nMobile\nSidebar becomes menu\nCards stack or reduce columns\nButtons wrap', 7.6, 2.9, 3.9, 2.0, 17, INK)
note(slide, 'The interface is designed as an operations tool: scanable tables, direct actions, and low-stock emphasis.')

# 15
slide = add_slide('Run it', 'How to start the application locally', 15)
code_box(slide, "# Terminal 1 - backend\ncd backend\npython -m venv .venv\n.venv\\Scripts\\Activate.ps1\npython -m pip install -r requirements.txt\npython app.py", 0.75, 1.85, 5.8, 2.65, 13)
code_box(slide, "# Terminal 2 - frontend\ncd frontend\nnpm install\nnpm run dev", 6.8, 1.85, 4.9, 1.65, 13)
rect(slide, 0.9, 5.05, 10.8, 0.78, PALE_GREEN, PALE_GREEN, radius=True)
textbox(slide, 'Backend: http://localhost:5000    Frontend: usually http://localhost:5173', 1.15, 5.28, 10.3, 0.28, 17, GREEN, True, align=PP_ALIGN.CENTER)
textbox(slide, 'Health check: http://localhost:5000/api/health', 3.2, 6.1, 6.5, 0.28, 15, BLUE, True, MONO, align=PP_ALIGN.CENTER)
note(slide, 'Both processes must be running: the browser UI depends on the Flask API for login and data.')

# 16
slide = add_slide('Demo', 'A complete demonstration sequence', 16)
steps = ['Start Flask and Vite', 'Log in as nishi / admin', 'Open Users and add a staff account', 'Add a supplier', 'Add a product with opening stock', 'Record a purchase', 'Create a customer bill', 'Open Reports and export history']
for i, text in enumerate(steps):
    x = 0.8 + (i % 2) * 6.0; y = 1.82 + (i // 2) * 1.08
    rect(slide, x, y, 5.35, 0.72, WHITE, RGBColor(225, 234, 227), radius=True)
    rect(slide, x + 0.14, y + 0.13, 0.45, 0.45, GREEN, GREEN, radius=True)
    textbox(slide, str(i + 1), x + 0.14, y + 0.2, 0.45, 0.22, 12, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, text, x + 0.78, y + 0.2, 4.35, 0.28, 15, INK, True)
note(slide, 'This sequence proves the important state changes: users, stock increases, stock decreases, invoices, and reports.')

# 17
slide = add_slide('Current limits', 'Important development notes', 17)
bullet_box(slide, ['Storage is in memory: restarting Flask removes products, suppliers, purchases, sales, and newly created users.', 'database/seed_data.json is reference data only and is not loaded at startup.', 'The CSV download is minimal: it serializes object values and does not add robust quoting or column headers.', 'The API URL is hard-coded to localhost:5000, so deployment needs environment-based configuration.', 'The README account table is out of sync with the actual code: code uses liki and aliya as staff accounts.'], 0.7, 1.85, 6.15, 4.45, 16)
rect(slide, 7.35, 1.95, 4.65, 3.9, PALE_YELLOW, PALE_YELLOW, radius=True)
textbox(slide, 'Recommended next steps', 7.7, 2.28, 3.95, 0.35, 19, YELLOW, True)
textbox(slide, '1. Add MySQL repository layer\n2. Move API URL to environment config\n3. Add automated API and UI tests\n4. Improve CSV generation\n5. Align README and demo accounts\n6. Add production token/session strategy', 7.7, 2.9, 3.75, 2.25, 16, INK)
note(slide, 'These are architecture and production-readiness concerns, not blockers for the local development demo.')

# 18
slide = add_slide('Recap', 'The complete mental model', 18, DARK)
textbox(slide, 'Browser', 0.9, 2.0, 2.2, 0.45, 23, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 'state + render + events', 0.85, 2.65, 2.3, 0.35, 15, RGBColor(211, 226, 216), align=PP_ALIGN.CENTER)
line(slide, 3.2, 2.45, 4.15, 2.45, RGBColor(143, 210, 170), 2)
textbox(slide, 'Flask API', 4.35, 2.0, 2.2, 0.45, 23, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 'auth + rules + JSON', 4.3, 2.65, 2.3, 0.35, 15, RGBColor(211, 226, 216), align=PP_ALIGN.CENTER)
line(slide, 6.65, 2.45, 7.6, 2.45, RGBColor(143, 210, 170), 2)
textbox(slide, 'Records', 7.8, 2.0, 2.2, 0.45, 23, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 'in-memory today, MySQL later', 7.65, 2.65, 2.55, 0.35, 15, RGBColor(211, 226, 216), align=PP_ALIGN.CENTER)
textbox(slide, 'UI action -> API request -> validated change -> refreshed state -> new screen', 1.35, 4.25, 10.6, 0.6, 23, RGBColor(143, 210, 170), True, align=PP_ALIGN.CENTER)
textbox(slide, 'That loop explains nearly every feature in the application.', 2.3, 5.2, 8.7, 0.35, 17, WHITE, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
