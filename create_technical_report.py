from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt

OUT = "technical_report_grocery_inventory.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(247, 248, 244)
INK = RGBColor(28, 37, 34)
MUTED = RGBColor(91, 105, 98)
GREEN = RGBColor(41, 116, 91)
LIME = RGBColor(184, 218, 94)
BLUE = RGBColor(55, 111, 145)
AMBER = RGBColor(202, 130, 45)
RED = RGBColor(177, 70, 62)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(216, 223, 216)


def add_bg(slide, section="TECHNICAL REPORT"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.16))
    bar.fill.solid(); bar.fill.fore_color.rgb = GREEN; bar.line.fill.background()
    tx = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(7), Inches(0.25))
    p = tx.text_frame.paragraphs[0]; p.text = section; p.font.size = Pt(9); p.font.bold = True; p.font.color.rgb = GREEN


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return box


def title(slide, heading, sub=None):
    textbox(slide, heading, 0.55, 0.72, 12.1, 0.55, 27, INK, True)
    if sub: textbox(slide, sub, 0.58, 1.34, 11.8, 0.35, 11, MUTED)


def bullets(slide, items, x, y, w, h, size=16, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {item}"; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color; p.space_after = Pt(10)
    return box


def card(slide, x, y, w, h, heading, body, accent=GREEN, body_size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = WHITE; shape.line.color.rgb = LINE
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent; strip.line.fill.background()
    textbox(slide, heading, x+0.25, y+0.2, w-0.45, 0.3, 15, INK, True)
    textbox(slide, body, x+0.25, y+0.62, w-0.45, h-0.78, body_size, MUTED)


def footer(slide, n):
    textbox(slide, f"Grocery Inventory Management  |  23 Aug 2026  |  {n}", 0.58, 7.16, 12.1, 0.18, 8, MUTED)


def new_slide(section="TECHNICAL REPORT"):
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s, section); return s

# 1 Cover
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s, "ENGINEERING REVIEW")
textbox(s, "Grocery Inventory\nManagement", 0.75, 1.35, 8.8, 1.35, 38, INK, True)
textbox(s, "Technical architecture and delivery report", 0.8, 2.9, 7, 0.4, 19, GREEN, True)
textbox(s, "A code-grounded review of the React/Vite frontend, Flask REST API,\ncurrent runtime behavior, risks, and recommended path to production.", 0.8, 3.55, 7.5, 0.75, 15, MUTED)
card(s, 9.25, 1.55, 3.15, 1.2, "STACK", "React + Vite\nFlask + Flask-CORS", BLUE)
card(s, 9.25, 3.05, 3.15, 1.2, "CURRENT STATE", "Functional prototype\nIn-memory persistence", AMBER)
card(s, 9.25, 4.55, 3.15, 1.2, "REVIEW DATE", "23 August 2026\nSource snapshot", GREEN)
footer(s, 1)

# 2 executive summary
s = new_slide("EXECUTIVE SUMMARY"); title(s, "The product works as a focused prototype", "The core inventory workflow is present; production readiness is limited by state, validation, and operational controls.")
card(s, 0.65, 2.0, 3.85, 2.15, "WHAT IS STRONG", "Clear user journey from login to inventory, purchasing, billing, and reporting. Server-side role checks exist, and password hashes plus random bearer tokens are used.", GREEN)
card(s, 4.75, 2.0, 3.85, 2.15, "PRIMARY CONSTRAINT", "All operational collections are process-local Python lists. Seed JSON is reference-only and is not loaded. A restart erases products, suppliers, purchases, and sales.", AMBER)
card(s, 8.85, 2.0, 3.85, 2.15, "RELEASE SIGNAL", "Frontend production build succeeds. No automated tests were found. The next investment should be a persistence and contract-testing foundation before feature expansion.", RED)
textbox(s, "Overall assessment", 0.75, 4.75, 2.2, 0.3, 14, MUTED, True)
textbox(s, "Prototype-ready  →  Not yet production-ready", 0.75, 5.18, 7.4, 0.55, 25, INK, True)
textbox(s, "Confidence: high for structure and observed code paths; runtime behavior should be confirmed with endpoint smoke tests and browser checks.", 0.78, 5.95, 11.3, 0.35, 11, MUTED)
footer(s, 2)

# 3 architecture
s = new_slide("ARCHITECTURE"); title(s, "Simple boundary today, thin domain layer", "The API boundary is the right migration seam, but business rules currently live inside one Flask module.")
# boxes
card(s, 0.7, 2.1, 3.1, 2.2, "BROWSER", "React single-page app\nApp owns auth, page state, loading, modals\nlocalStorage holds token + user", BLUE)
card(s, 5.1, 2.1, 3.1, 2.2, "HTTP API", "Flask routes under /api\nBearer auth decorator\nCORS enabled globally", GREEN)
card(s, 9.5, 2.1, 3.1, 2.2, "RUNTIME STATE", "Python lists: products, suppliers, purchases, sales\nPython dict: users, active_tokens", AMBER)
for x in (3.95, 8.35):
    arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(2.85), Inches(0.85), Inches(0.55)); arrow.fill.solid(); arrow.fill.fore_color.rgb = LIME; arrow.line.fill.background()
textbox(s, "Future target: route → service/repository → MySQL", 3.7, 5.25, 5.9, 0.4, 20, GREEN, True, PP_ALIGN.CENTER)
textbox(s, "Verified files: frontend/src/main.jsx, frontend/vite.config.js, backend/app.py, database/seed_data.json", 1.0, 6.1, 11.2, 0.3, 11, MUTED, False, PP_ALIGN.CENTER)
footer(s, 3)

# 4 runtime flow
s = new_slide("RUNTIME FLOW"); title(s, "Authentication and data loading are centralized in App", "One load cycle fans out to the summary and operational collections; failures are tolerated individually.")
steps = [("01", "Login", "POST /api/auth/login"), ("02", "Persist session", "localStorage token + user"), ("03", "Load workspace", "Promise.allSettled across 6 requests"), ("04", "Mutate stock", "Purchase adds; sale subtracts"), ("05", "Refresh UI", "onSaved / onSold calls loadData")]
for i, (num, head, body) in enumerate(steps):
    x = 0.7 + i*2.48
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(2.4), Inches(0.62), Inches(0.62)); circle.fill.solid(); circle.fill.fore_color.rgb = GREEN; circle.line.fill.background()
    textbox(s, num, x, 2.56, 0.62, 0.18, 11, WHITE, True, PP_ALIGN.CENTER)
    textbox(s, head, x, 3.25, 2.0, 0.3, 15, INK, True)
    textbox(s, body, x, 3.72, 2.0, 0.65, 12, MUTED)
    if i < 4:
        line = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x+1.55), Inches(2.56), Inches(0.68), Inches(0.25)); line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
card(s, 1.0, 5.1, 5.3, 1.0, "GOOD", "One request helper adds Authorization consistently and handles 401 expiry centrally.", GREEN, 12)
card(s, 6.95, 5.1, 5.3, 1.0, "WATCH", "Promise.allSettled can leave partial or stale screens without surfacing which endpoint failed.", RED, 12)
footer(s, 4)

# 5 domain and API
s = new_slide("DOMAIN AND API"); title(s, "The API covers the core operating model", "Routes map cleanly to inventory, suppliers, purchasing, sales, users, and summary reporting.")
rows = [
    ("Auth", "POST /auth/login  |  GET /auth/me", "Bearer token; all protected routes"),
    ("Catalog", "GET /products  |  POST/PATCH /products", "Admin writes; authenticated reads"),
    ("Operations", "POST /purchases  |  POST /sales", "Stock and supplier quantities mutate"),
    ("Reference", "GET /suppliers  |  GET /categories", "Supplier creation is admin-only"),
    ("Reporting", "GET /summary  |  GET /sales", "Summary includes today totals and profit"),
    ("Administration", "CRUD /users", "Admin-only user lifecycle"),
]
# table
x0, y0 = 0.75, 1.95
widths = [1.7, 4.3, 5.75]
headers = ["AREA", "ROUTES", "POLICY / EFFECT"]
for j, h in enumerate(headers):
    textbox(s, h, x0 + sum(widths[:j]), y0, widths[j], 0.28, 10, GREEN, True)
for i, row in enumerate(rows):
    y = y0 + 0.48 + i*0.68
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y-0.08), Inches(sum(widths)), Inches(0.56)); band.fill.solid(); band.fill.fore_color.rgb = WHITE if i%2==0 else RGBColor(238,243,237); band.line.fill.background()
    for j, value in enumerate(row): textbox(s, value, x0 + sum(widths[:j])+0.12, y, widths[j]-0.2, 0.34, 12, INK if j==0 else MUTED, j==0)
textbox(s, "Key contract gap: input validation is distributed inside handlers, with several direct int(...) conversions that can become 500 errors for malformed payloads.", 0.8, 6.45, 11.8, 0.4, 12, RED, True)
footer(s, 5)

# 6 data model
s = new_slide("DATA AND PERSISTENCE"); title(s, "State is usable for demos, fragile for operations", "The code has a workable conceptual model, but no durable source of truth or transactional boundary.")
card(s, 0.7, 1.95, 2.75, 2.25, "PRODUCT", "id, sku, name, category, price, quantity, reorder_level, unit, supplier", GREEN)
card(s, 3.75, 1.95, 2.75, 2.25, "SUPPLIER", "id, name, contact, products, supply_quantities, supply_prices, status", BLUE)
card(s, 6.8, 1.95, 2.75, 2.25, "PURCHASE", "id, supplier, products, total, date, status; increases stock", AMBER)
card(s, 9.85, 1.95, 2.75, 2.25, "SALE", "id, products, total, date; decreases stock", RED)
textbox(s, "Current behavior", 0.75, 4.75, 2.2, 0.3, 14, MUTED, True)
bullets(s, ["Collections initialize empty in backend/app.py.", "database/seed_data.json is not loaded at startup.", "No migrations, repository, transaction, or concurrency protection.", "Profit today is revenue minus purchase totals, not COGS-based margin."], 0.82, 5.1, 7.0, 1.5, 14)
textbox(s, "Migration implication", 8.25, 4.75, 2.2, 0.3, 14, MUTED, True)
textbox(s, "Introduce a repository/service layer before MySQL. Preserve route contracts while moving stock adjustments into atomic transactions.", 8.25, 5.12, 4.2, 1.1, 16, INK, True)
footer(s, 6)

# 7 security
s = new_slide("SECURITY REVIEW"); title(s, "Baseline controls exist; deployment controls do not", "Authentication is present, but session lifecycle and production configuration need hardening.")
card(s, 0.7, 1.9, 3.8, 1.45, "POSITIVE CONTROLS", "Werkzeug password hashing\nsecrets.token_urlsafe bearer tokens\nserver-side role enforcement", GREEN)
card(s, 4.75, 1.9, 3.8, 1.45, "HIGH PRIORITY", "Non-expiring in-memory tokens\nGlobal open CORS\nFlask debug=True", RED)
card(s, 8.8, 1.9, 3.8, 1.45, "CLIENT EXPOSURE", "Bearer token in localStorage\nDemo credentials shown in login UI and README", AMBER)
bullets(s, ["Add token expiry and revocation semantics backed by durable session storage.", "Restrict CORS by environment and use production secret/config management.", "Disable debug mode outside local development.", "Move demo accounts and credentials behind an explicit development-only configuration."], 0.9, 4.35, 11.5, 1.6, 16)
footer(s, 7)

# 8 quality
s = new_slide("QUALITY AND DELIVERY"); title(s, "Build confidence is ahead of behavior confidence", "The frontend compiles, but the project has no automated safety net for its most important workflows.")
card(s, 0.8, 1.95, 3.55, 2.0, "BUILD", "PASS\nVite production build completes successfully.\n1815 modules transformed.", GREEN, 14)
card(s, 4.9, 1.95, 3.55, 2.0, "TESTS", "GAP\nNo test files or test scripts found.\nNo endpoint smoke suite.", RED, 14)
card(s, 9.0, 1.95, 3.55, 2.0, "DIAGNOSTICS", "OBSERVED\nNo source diagnostics reported in the reviewed JSX, CSS, or Python surfaces.", BLUE, 14)
textbox(s, "Minimum regression suite", 0.85, 4.55, 3.4, 0.3, 15, INK, True)
bullets(s, ["Login success, invalid credentials, and role restrictions.", "Purchase stock increase and supplier quantity decrement.", "Sale stock decrement and insufficient-stock rejection.", "Malformed numeric payloads return controlled 4xx responses.", "Summary totals reflect the intended accounting definition."], 0.9, 4.95, 11.2, 1.7, 14)
footer(s, 8)

# 9 risks
s = new_slide("RISKS AND TECHNICAL DEBT"); title(s, "Five risks should shape the next iteration", "Prioritized by data integrity and operational impact.")
risks = [
    ("P0", "Data loss on restart", "No persistence; every operational record is process-local.", RED),
    ("P1", "Non-atomic stock updates", "Purchase and sale mutations have no database transaction or concurrency guard.", RED),
    ("P1", "Weak request contracts", "Handler-level validation and direct conversions can produce server errors.", AMBER),
    ("P1", "Security defaults", "Debug mode, open CORS, non-expiring tokens, and localStorage bearer tokens.", AMBER),
    ("P2", "Product consistency", "Supplier counts/reassignment and category data can drift; README permissions differ from UI.", BLUE),
]
for i, (prio, head, body, accent) in enumerate(risks):
    y = 1.85 + i*0.92
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(0.78), Inches(0.45)); tag.fill.solid(); tag.fill.fore_color.rgb = accent; tag.line.fill.background()
    textbox(s, prio, 0.8, y+0.12, 0.78, 0.18, 11, WHITE, True, PP_ALIGN.CENTER)
    textbox(s, head, 1.85, y+0.02, 3.0, 0.3, 15, INK, True)
    textbox(s, body, 5.0, y+0.03, 7.4, 0.35, 13, MUTED)
footer(s, 9)

# 10 roadmap
s = new_slide("RECOMMENDED ROADMAP"); title(s, "Sequence the work around a durable core", "The goal is to preserve the current user experience while making state and contracts trustworthy.")
road = [
    ("1", "Stabilize contracts", "Add schema validation, consistent error envelopes, endpoint tests, and a development seed loader.", GREEN),
    ("2", "Extract domain services", "Move stock adjustments, supplier logic, and summary calculations out of route handlers.", BLUE),
    ("3", "Add persistence", "Introduce MySQL schema, migrations, repository interfaces, and atomic inventory transactions.", AMBER),
    ("4", "Harden deployment", "Environment-based API URL, restricted CORS, token expiry, secrets, and debug off.", RED),
    ("5", "Improve product truth", "COGS-based profit, audit history, pagination/filtering, and UI/API permission alignment.", GREEN),
]
for i, (num, head, body, accent) in enumerate(road):
    x = 0.75 + i*2.48
    top = s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(2.0), Inches(0.7), Inches(0.7)); top.fill.solid(); top.fill.fore_color.rgb = accent; top.line.fill.background()
    textbox(s, num, x, 2.22, 0.7, 0.2, 16, WHITE, True, PP_ALIGN.CENTER)
    textbox(s, head, x, 3.05, 2.15, 0.4, 14, INK, True)
    textbox(s, body, x, 3.62, 2.12, 1.5, 12, MUTED)
    if i < 4:
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x+1.93), Inches(2.22), Inches(0.46), Inches(0.2)); arrow.fill.solid(); arrow.fill.fore_color.rgb = LINE; arrow.line.fill.background()
textbox(s, "Decision point: invest in persistence and contracts before adding more screens or reports.", 0.85, 6.15, 11.8, 0.4, 18, GREEN, True, PP_ALIGN.CENTER)
footer(s, 10)

# 11 appendix / sources
s = new_slide("APPENDIX"); title(s, "Source map and review boundaries", "This deck is grounded in the repository snapshot available on 23 August 2026.")
card(s, 0.75, 1.9, 5.8, 2.0, "PRIMARY SOURCES", "backend/app.py\nfrontend/src/main.jsx\nfrontend/src/styles.css\nfrontend/package.json\nbackend/requirements.txt", GREEN, 14)
card(s, 6.8, 1.9, 5.8, 2.0, "DATA / OPERATIONS", "database/seed_data.json\nREADME.md\nVite production build: PASS\nNo automated tests found", BLUE, 14)
textbox(s, "Review boundary", 0.8, 4.65, 2.0, 0.3, 15, INK, True)
textbox(s, "This is a static code and configuration assessment. It does not claim production readiness, penetration-test coverage, database performance, or browser-level UX verification.", 0.8, 5.05, 11.5, 0.75, 17, MUTED)
textbox(s, "End of report", 0.8, 6.25, 11.5, 0.4, 22, GREEN, True, PP_ALIGN.CENTER)
footer(s, 11)

prs.save(OUT)
print(OUT)
